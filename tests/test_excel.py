"""Check that the workbook's formulas agree with the Python model.

Needs LibreOffice (``soffice``) to recalculate the formulas; skipped otherwise.
"""

import shutil
import subprocess
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from installment_tracker.excel import FIRST_ROW, build_workbook
from installment_tracker.model import Purchase, build_schedule

soffice = shutil.which("soffice") or shutil.which("libreoffice")
pytestmark = pytest.mark.skipif(soffice is None, reason="LibreOffice not installed")

PURCHASES = [
    Purchase("P1", date(2026, 5, 20), "Supplier A", 1664.40, 4),
    Purchase("P2", date(2026, 6, 11), "Supplier B", 2548.77, 6),
    Purchase("P3", date(2026, 6, 30), "Supplier B", 4094.26, 6, first_month_override=date(2026, 9, 1)),
    Purchase("P4", date(2026, 8, 14), "", 914.68, 1),
]


def _recalculate(path: Path) -> Path:
    subprocess.run(
        [soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(path.parent / "out"), str(path)],
        check=True, capture_output=True, timeout=120,
    )
    return path.parent / "out" / path.name


def test_formulas_match_model(config, tmp_path):
    xlsx = build_workbook(PURCHASES, config, tmp_path / "t.xlsx")
    ws = load_workbook(_recalculate(xlsx), data_only=True)["Purchases"]
    today = date.today()
    for i, p in enumerate(PURCHASES):
        r = FIRST_ROW + i
        s = build_schedule(p, config)
        assert ws[f"H{r}"].value == pytest.approx(s.commission)
        assert ws[f"J{r}"].value == pytest.approx(s.monthly_installment)
        assert ws[f"K{r}"].value == pytest.approx(s.first_installment)
        assert ws[f"L{r}"].value.date() == s.first_month
        assert ws[f"M{r}"].value.date() == s.last_month
        assert ws[f"N{r}"].value == s.installments_paid(today, config.payment_day)
        assert ws[f"P{r}"].value == pytest.approx(s.outstanding(today, config.payment_day))
        assert ws[f"Q{r}"].value == s.status(today, config.payment_day)


def test_calendar_month_totals(config, tmp_path):
    xlsx = build_workbook(PURCHASES, config, tmp_path / "t.xlsx")
    wb = load_workbook(_recalculate(xlsx), data_only=True)
    ws = wb["Calendar"]
    totals_row = FIRST_ROW + len(PURCHASES) + config.spare_rows + 1
    for col in range(4, 4 + config.calendar_months):
        month = ws.cell(row=4, column=col).value.date()
        expected = sum(build_schedule(p, config).installment_for(month) for p in PURCHASES)
        assert ws.cell(row=totals_row, column=col).value == pytest.approx(expected)
