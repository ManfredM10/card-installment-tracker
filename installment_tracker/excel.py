"""Build the tracker workbook with live formulas.

The workbook is not a static report: every derived value is an Excel formula
that references the ``Settings`` sheet, so the file keeps working after the
user adds purchases by hand or changes a commission rate. The Python model in
:mod:`installment_tracker.model` implements the same rules and is used by the
tests to check the formulas' output.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .config import Config
from .model import Purchase

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF", size=10)       # user inputs
BLACK = Font(name=FONT, color="000000", size=10)      # formulas
GREEN = Font(name=FONT, color="008000", size=10)      # links to another sheet
BOLD = Font(name=FONT, bold=True, size=10)
TITLE = Font(name=FONT, bold=True, size=14)
NOTE = Font(name=FONT, italic=True, size=9, color="595959")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
FILL_PAID = PatternFill("solid", fgColor="C6EFCE")
FILL_DUE = PatternFill("solid", fgColor="FFEB9C")
FILL_PENDING = PatternFill("solid", fgColor="DDEBF7")
_thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
PCT = "0.00%"
DATEF = "dd/mm/yyyy"
MONF = "mmm yyyy"

HEADER_ROW = 4
FIRST_ROW = 5

# Settings sheet cell addresses (absolute, used by every formula)
S_TODAY = "Settings!$B$4"
S_FX = "Settings!$B$5"
S_PAYDAY = "Settings!$B$6"
S_OFFSET = "Settings!$B$7"
COMMISSION_TABLE_ROW = 12


def _money(symbol: str, decimals: int = 2) -> str:
    dec = "." + "0" * decimals if decimals else ""
    return f'"{symbol}"#,##0{dec};("{symbol}"#,##0{dec});"-"'


def _header(ws, row, headers, widths=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(purchases: list[Purchase], config: Config, output: str | Path) -> Path:
    """Write the workbook to ``output`` and return its path."""
    usd = _money(config.currency_symbol)
    has_secondary = bool(config.secondary_currency_symbol)
    sec = _money(config.secondary_currency_symbol, 0) if has_secondary else usd

    n_rows = len(purchases) + config.spare_rows
    last_row = FIRST_ROW + n_rows - 1
    totals_row = last_row + 2
    n_terms = len(config.commissions)
    terms_range = f"Settings!$A${COMMISSION_TABLE_ROW}:$A${COMMISSION_TABLE_ROW + n_terms - 1}"
    rates_range = f"Settings!$B${COMMISSION_TABLE_ROW}:$B${COMMISSION_TABLE_ROW + n_terms - 1}"
    merchants_first = COMMISSION_TABLE_ROW + n_terms + 3

    wb = Workbook()
    _build_settings(wb.active, config, merchants_first)
    _build_purchases(
        wb.create_sheet("Purchases"), purchases, config, last_row, totals_row,
        terms_range, rates_range, merchants_first, usd, sec, has_secondary,
    )
    calendar_start = _earliest_month(purchases)
    _build_calendar(wb.create_sheet("Calendar"), config, last_row, totals_row, calendar_start, usd, sec, has_secondary)
    _build_summary(wb.create_sheet("Summary", 0), config, last_row, totals_row, merchants_first, usd, sec, has_secondary)

    out = Path(output)
    wb.save(out)
    return out


def _earliest_month(purchases: list[Purchase]) -> tuple[int, int]:
    if not purchases:
        from datetime import date
        t = date.today()
        return t.year, t.month
    d = min(p.purchase_date for p in purchases)
    return d.year, d.month


def _build_settings(ws, config: Config, merchants_first: int) -> None:
    ws.title = "Settings"
    ws["A1"] = "Settings"; ws["A1"].font = TITLE
    ws["A2"] = "Blue cells on yellow are editable. Everything else in the workbook is calculated."; ws["A2"].font = NOTE

    rows = [
        ("Today", "=TODAY()", DATEF, False, "Updates automatically. Decides which installments count as paid."),
        (f"Exchange rate ({config.secondary_currency_symbol or 'n/a'} per {config.currency_symbol})",
         config.exchange_rate, "#,##0.00", True,
         'In Google Sheets you can replace this with =GOOGLEFINANCE("CURRENCY:USDCRC") (adjust the pair).'),
        ("Card payment day", config.payment_day, "0", True,
         "An installment is considered paid once this day of its month has passed."),
        ("Months until first installment", config.first_installment_offset, "0", True,
         "Purchase in May with 2 → first installment in July."),
    ]
    for i, (label, value, fmt, editable, note) in enumerate(rows):
        r = 4 + i
        ws.cell(row=r, column=1, value=label).font = BOLD
        c = ws.cell(row=r, column=2, value=value)
        c.number_format, c.border = fmt, BORDER
        c.font = BLUE if editable else BLACK
        if editable:
            c.fill = INPUT_FILL
        ws.cell(row=r, column=3, value=note).font = NOTE

    ws["A10"] = "Commission by term"; ws["A10"].font = BOLD
    _header(ws, 11, ["Term (months)", "Commission %"])
    for i, term in enumerate(config.allowed_terms):
        r = COMMISSION_TABLE_ROW + i
        a = ws.cell(row=r, column=1, value=term)
        b = ws.cell(row=r, column=2, value=config.commissions[term])
        for c in (a, b):
            c.font, c.fill, c.border = BLUE, INPUT_FILL, BORDER
        b.number_format = PCT
    ws.cell(row=COMMISSION_TABLE_ROW, column=3, value="Term 1 = paid in full on the next statement.").font = NOTE

    ws.cell(row=merchants_first - 1, column=1, value="Merchants (dropdown list)").font = BOLD
    for i, m in enumerate(config.merchants):
        c = ws.cell(row=merchants_first + i, column=1, value=m)
        c.font, c.fill, c.border = BLUE, INPUT_FILL, BORDER
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 90


def _build_purchases(ws, purchases, config, last_row, totals_row, terms_range, rates_range,
                     merchants_first, usd, sec, has_secondary) -> None:
    sym = config.currency_symbol
    ws["A1"] = "Purchases"; ws["A1"].font = TITLE
    ws["A2"] = ("Type only the yellow columns (A–F). Commission, installments, dates, balance and status are "
                "formulas driven by the Settings sheet. 'First month (override)' is optional: use it only when the "
                "bank charged the first installment in a different month than usual.")
    ws["A2"].font = NOTE
    ws.merge_cells("A2:S2")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 30

    headers = ["Code", "Purchase date", "Merchant", f"Amount ({sym})", "Term (months)", "First month (override)",
               "Commission %", f"Commission ({sym})", f"Total financed ({sym})", f"Monthly installment ({sym})",
               f"First installment ({sym})", "First month", "Last month", "Installments paid",
               "Installments left", f"Outstanding ({sym})", "Status",
               f"Monthly installment ({config.secondary_currency_symbol})" if has_secondary else "Monthly (converted)",
               "Next installment"]
    widths = [10, 12, 12, 13, 9, 13, 10, 12, 14, 13, 13, 11, 11, 10, 10, 14, 10, 14, 12]
    _header(ws, HEADER_ROW, headers, widths)

    for r in range(FIRST_ROW, last_row + 1):
        R = str(r)
        formulas = {
            "G": f'=IF($A{R}="","",IFERROR(INDEX({rates_range},MATCH($E{R},{terms_range},0)),0))',
            "H": f'=IF($A{R}="","",ROUND($D{R}*$G{R},2))',
            "I": f'=IF($A{R}="","",$D{R}+$H{R})',
            "J": f'=IF($A{R}="","",ROUND($D{R}/$E{R},2))',
            "K": f'=IF($A{R}="","",$J{R}+$H{R})',
            "L": f'=IF($A{R}="","",IF($F{R}<>"",DATE(YEAR($F{R}),MONTH($F{R}),1),EOMONTH($B{R},{S_OFFSET}-1)+1))',
            "M": f'=IF($A{R}="","",EDATE($L{R},$E{R}-1))',
            "N": (f'=IF($A{R}="","",MAX(0,MIN($E{R},(YEAR({S_TODAY})-YEAR($L{R}))*12'
                  f'+MONTH({S_TODAY})-MONTH($L{R})+IF(DAY({S_TODAY})>={S_PAYDAY},1,0))))'),
            "O": f'=IF($A{R}="","",$E{R}-$N{R})',
            "P": f'=IF($A{R}="","",IF($N{R}=0,$I{R},$O{R}*$J{R}))',
            "Q": f'=IF($A{R}="","",IF($O{R}=0,"Paid",IF($N{R}=0,"Pending","Active")))',
            "R": f'=IF($A{R}="","",ROUND($J{R}*{S_FX},{0 if has_secondary else 2}))',
            "S": f'=IF($A{R}="","",IF($O{R}=0,"-",EDATE($L{R},$N{R})))',
        }
        for col in "ABCDEF":
            c = ws[f"{col}{R}"]
            c.font, c.fill, c.border = BLUE, INPUT_FILL, BORDER
        for col, f in formulas.items():
            c = ws[f"{col}{R}"]
            c.value, c.font, c.border = f, BLACK, BORDER
        ws[f"B{R}"].number_format = ws[f"F{R}"].number_format = DATEF
        for col in "DHIJKP":
            ws[f"{col}{R}"].number_format = usd
        ws[f"G{R}"].number_format = PCT
        for col in "LMS":
            ws[f"{col}{R}"].number_format = MONF
        ws[f"R{R}"].number_format = sec
        ws[f"Q{R}"].alignment = Alignment(horizontal="center")

    for i, p in enumerate(purchases):
        r = FIRST_ROW + i
        ws[f"A{r}"] = p.code
        ws[f"B{r}"] = p.purchase_date
        ws[f"C{r}"] = p.merchant or None
        ws[f"D{r}"] = p.amount
        ws[f"E{r}"] = p.term_months
        if p.first_month_override:
            ws[f"F{r}"] = p.first_month_override

    ws[f"A{totals_row}"] = "TOTALS"; ws[f"A{totals_row}"].font = BOLD
    for col in "DHIP":
        c = ws[f"{col}{totals_row}"]
        c.value, c.font, c.number_format = f"=SUM({col}{FIRST_ROW}:{col}{last_row})", BOLD, usd
    ws[f"O{totals_row}"] = "Total outstanding →"; ws[f"O{totals_row}"].font = BOLD

    n_merchants = max(1, len(config.merchants))
    merchants_range = f"Settings!$A${merchants_first}:$A${merchants_first + n_merchants - 1}"
    validations = [
        (DataValidation(type="list", formula1=f"={merchants_range}", allow_blank=True), [f"C{FIRST_ROW}:C{last_row}"]),
        (DataValidation(type="list", formula1=f"={terms_range}", allow_blank=True, showErrorMessage=True,
                        error="Term not in the commission table (see Settings)"), [f"E{FIRST_ROW}:E{last_row}"]),
        (DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True, showErrorMessage=True,
                        error="Type the number only, no currency symbols"), [f"D{FIRST_ROW}:D{last_row}"]),
        (DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, showErrorMessage=True,
                        error="Invalid date"), [f"B{FIRST_ROW}:B{last_row}", f"F{FIRST_ROW}:F{last_row}"]),
    ]
    for dv, ranges in validations:
        ws.add_data_validation(dv)
        for rng in ranges:
            dv.add(rng)

    status_rng = f"Q{FIRST_ROW}:Q{last_row}"
    for label, fill, color in (("Paid", FILL_PAID, "006100"), ("Active", FILL_DUE, "9C5700"), ("Pending", FILL_PENDING, "1F3864")):
        ws.conditional_formatting.add(
            status_rng, FormulaRule(formula=[f'$Q{FIRST_ROW}="{label}"'], fill=fill, font=Font(color=color)))
    ws.freeze_panes = "C5"


def _build_calendar(ws, config, last_row, totals_row, start, usd, sec, has_secondary) -> None:
    ws["A1"] = "Installment calendar (generated from Purchases)"; ws["A1"].font = TITLE
    ws["A2"] = ("Green = already paid · Yellow = due this month · No colour = pending. "
                "Row 3 shows each month's state given today's date and the card payment day.")
    ws["A2"].font = NOTE
    first_col = 4
    ws["C3"] = "Month state →"; ws["C3"].font = NOTE; ws["C3"].alignment = Alignment(horizontal="right")
    _header(ws, HEADER_ROW, ["Code", f"Total financed ({config.currency_symbol})", "Status"], [11, 14, 10])

    year, month = start
    for j in range(config.calendar_months):
        col = get_column_letter(first_col + j)
        h = ws[f"{col}{HEADER_ROW}"]
        h.value = f"=DATE({year},{month},1)" if j == 0 else f"=EDATE({get_column_letter(first_col + j - 1)}{HEADER_ROW},1)"
        h.number_format, h.font, h.fill, h.border = MONF, HDR_FONT, HDR_FILL, BORDER
        h.alignment = Alignment(horizontal="center")
        st = ws[f"{col}3"]
        st.value = (f'=IF(DATE(YEAR({col}$4),MONTH({col}$4),{S_PAYDAY})<={S_TODAY},"Paid",'
                    f'IF(AND(YEAR({col}$4)=YEAR({S_TODAY}),MONTH({col}$4)=MONTH({S_TODAY})),"Due","Pending"))')
        st.font = Font(name=FONT, italic=True, size=8, color="595959")
        st.alignment = Alignment(horizontal="center")
        ws.column_dimensions[col].width = 11
    ws.row_dimensions[HEADER_ROW].height = 20

    for r in range(FIRST_ROW, last_row + 1):
        R = str(r)
        ws[f"A{R}"] = f'=IF(Purchases!$A{R}="","",Purchases!$A{R})'
        ws[f"B{R}"] = f'=IF(Purchases!$A{R}="","",Purchases!$I{R})'
        ws[f"C{R}"] = f'=IF(Purchases!$A{R}="","",Purchases!$Q{R})'
        ws[f"B{R}"].number_format = usd
        ws[f"C{R}"].alignment = Alignment(horizontal="center")
        for col in "ABC":
            ws[f"{col}{R}"].font, ws[f"{col}{R}"].border = GREEN, BORDER
        for j in range(config.calendar_months):
            col = get_column_letter(first_col + j)
            c = ws[f"{col}{R}"]
            c.value = (f'=IF($A{R}="","",IF({col}$4=Purchases!$L{R},Purchases!$K{R},'
                       f'IF(AND({col}$4>Purchases!$L{R},{col}$4<=Purchases!$M{R}),Purchases!$J{R},"")))')
            c.number_format, c.font, c.border = usd, BLACK, BORDER

    ws[f"A{totals_row}"] = f"MONTH TOTAL ({config.currency_symbol})"; ws[f"A{totals_row}"].font = BOLD
    if has_secondary:
        ws[f"A{totals_row + 1}"] = f"MONTH TOTAL ({config.secondary_currency_symbol})"
        ws[f"A{totals_row + 1}"].font = BOLD
    for j in range(config.calendar_months):
        col = get_column_letter(first_col + j)
        c = ws[f"{col}{totals_row}"]
        c.value, c.font, c.number_format = f"=SUM({col}{FIRST_ROW}:{col}{last_row})", BOLD, usd
        if has_secondary:
            c2 = ws[f"{col}{totals_row + 1}"]
            c2.value, c2.font, c2.number_format = f"=ROUND({col}{totals_row}*{S_FX},0)", BOLD, sec

    last_col = get_column_letter(first_col + config.calendar_months - 1)
    grid = f"D{FIRST_ROW}:{last_col}{last_row}"
    ws.conditional_formatting.add(grid, FormulaRule(formula=[f'AND(D{FIRST_ROW}<>"",D$3="Paid")'], fill=FILL_PAID))
    ws.conditional_formatting.add(grid, FormulaRule(formula=[f'AND(D{FIRST_ROW}<>"",D$3="Due")'], fill=FILL_DUE))
    status_rng = f"C{FIRST_ROW}:C{last_row}"
    ws.conditional_formatting.add(status_rng, FormulaRule(formula=[f'$C{FIRST_ROW}="Paid"'], fill=FILL_PAID))
    ws.conditional_formatting.add(status_rng, FormulaRule(formula=[f'$C{FIRST_ROW}="Active"'], fill=FILL_DUE))
    ws.freeze_panes = "D5"


def _build_summary(ws, config, last_row, totals_row, merchants_first, usd, sec, has_secondary) -> None:
    sym = config.currency_symbol
    last_col = get_column_letter(4 + config.calendar_months - 1)
    cal_hdr = f"Calendar!D{HEADER_ROW}:{last_col}{HEADER_ROW}"
    cal_tot = f"Calendar!D{totals_row}:{last_col}{totals_row}"
    P = lambda col: f"Purchases!{col}{FIRST_ROW}:{col}{last_row}"  # noqa: E731

    ws["A1"] = "Summary"; ws["A1"].font = TITLE
    ws["A2"] = "Everything here is calculated; nothing to type."; ws["A2"].font = NOTE
    rows = [
        ("Today", f"={S_TODAY}", DATEF),
        ("Purchases recorded", f"=COUNTA({P('A')})", "0"),
        ("Active", f'=COUNTIF({P("Q")},"Active")', "0"),
        ("Pending (not started)", f'=COUNTIF({P("Q")},"Pending")', "0"),
        ("Paid", f'=COUNTIF({P("Q")},"Paid")', "0"),
        (f"Total purchased ({sym})", f"=SUM({P('D')})", usd),
        (f"Commissions paid to the issuer ({sym})", f"=SUM({P('H')})", usd),
        (f"Total outstanding ({sym})", f"=SUM({P('P')})", usd),
        (f"This month's installments ({sym})",
         f"=SUMPRODUCT((YEAR({cal_hdr})=YEAR({S_TODAY}))*(MONTH({cal_hdr})=MONTH({S_TODAY}))*{cal_tot})", usd),
        (f"Next month's installments ({sym})",
         f"=SUMPRODUCT((YEAR({cal_hdr})=YEAR(EDATE({S_TODAY},1)))*(MONTH({cal_hdr})=MONTH(EDATE({S_TODAY},1)))*{cal_tot})", usd),
    ]
    if has_secondary:
        rows.insert(8, (f"Total outstanding ({config.secondary_currency_symbol})", f"=ROUND(SUM({P('P')})*{S_FX},0)", sec))
    for i, (label, f, fmt) in enumerate(rows):
        r = 4 + i
        ws.cell(row=r, column=1, value=label).font = BOLD
        c = ws.cell(row=r, column=2, value=f)
        c.number_format, c.font, c.border = fmt, GREEN, BORDER

    r0 = 4 + len(rows) + 2
    ws.cell(row=r0, column=1, value="By merchant").font = BOLD
    _header(ws, r0 + 1, ["Merchant", "Purchases", f"Total purchased ({sym})", f"Outstanding ({sym})"])
    n = len(config.merchants)
    for i in range(n):
        r = r0 + 2 + i
        ws.cell(row=r, column=1, value=f"=Settings!A{merchants_first + i}").font = GREEN
        ws.cell(row=r, column=2, value=f"=COUNTIF({P('C')},A{r})").number_format = "0"
        ws.cell(row=r, column=3, value=f"=SUMIF({P('C')},A{r},{P('D')})").number_format = usd
        ws.cell(row=r, column=4, value=f"=SUMIF({P('C')},A{r},{P('P')})").number_format = usd
        for col in range(1, 5):
            ws.cell(row=r, column=col).border = BORDER
    r = r0 + 2 + n
    rng_c = f"C{r0 + 2}:C{r - 1}" if n else "0"
    rng_d = f"D{r0 + 2}:D{r - 1}" if n else "0"
    ws.cell(row=r, column=1, value="No merchant assigned").font = BOLD
    ws.cell(row=r, column=2, value=f"=COUNTA({P('A')})-COUNTA({P('C')})").number_format = "0"
    ws.cell(row=r, column=3, value=f"=SUM({P('D')})-SUM({rng_c})").number_format = usd
    ws.cell(row=r, column=4, value=f"=SUM({P('P')})-SUM({rng_d})").number_format = usd
    for col in range(1, 5):
        ws.cell(row=r, column=col).border = BORDER
    for col, w in zip("ABCD", (36, 16, 18, 18)):
        ws.column_dimensions[col].width = w
